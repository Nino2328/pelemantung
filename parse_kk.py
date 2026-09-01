from collections import Counter
import json

from openpyxl import load_workbook

path = r'C:\Users\ADVAN\Downloads\data KK Pelemantung, Sidamesra thn 2022.xlsx'
out_path = r'C:\React\Pelemantung\pelemantung\kk_stats.json'

wb = load_workbook(path, read_only=True, data_only=True)
ws = wb[wb.sheetnames[0]]
rows = list(ws.iter_rows(values_only=True))
if not rows:
    raise ValueError('Spreadsheet is empty')

header = rows[0]
header_map = {name: idx for idx, name in enumerate(header)}

sex_idx = header_map.get('JENIS KELAMIN')
kk_idx = header_map.get('No.KK')
rt_idx = header_map.get('RT')

if sex_idx is None or kk_idx is None or rt_idx is None:
    raise ValueError('Required columns not found in spreadsheet')

records = rows[1:]
sex_counts = Counter()
unique_kks = set()
unique_rts = set()

for row in records:
    if row[sex_idx] is not None:
        sex_counts[str(row[sex_idx]).strip()] += 1
    if row[kk_idx] is not None:
        unique_kks.add(str(row[kk_idx]).strip())
    if row[rt_idx] is not None:
        unique_rts.add(str(row[rt_idx]).strip())

summary = {
    'total_population': len(records),
    'total_households': len(unique_kks),
    'male': sex_counts.get('Laki-laki', 0),
    'female': sex_counts.get('Perempuan', 0),
    'rt_count': len(unique_rts),
    'rt_values': sorted(unique_rts),
}

with open(out_path, 'w', encoding='utf-8') as f:
    json.dump(summary, f, ensure_ascii=False, indent=2)

print(json.dumps(summary, ensure_ascii=False, indent=2))
